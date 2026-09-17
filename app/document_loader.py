from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, List


SUPPORTED_SUFFIXES = {".md", ".txt", ".pdf", ".docx", ".xlsx"}


@dataclass(frozen=True)
class DocumentPart:
    source: str
    locator: str
    text: str


def display_path(path: Path) -> str:
    resolved = path.resolve()
    try:
        return resolved.relative_to(Path.cwd().resolve()).as_posix()
    except ValueError:
        return resolved.name


def iter_supported_files(inputs: Iterable[Path]) -> List[Path]:
    files: List[Path] = []
    for input_path in inputs:
        path = input_path.expanduser()
        if path.is_file():
            if path.suffix.lower() in SUPPORTED_SUFFIXES:
                files.append(path)
            continue
        if path.is_dir():
            files.extend(
                candidate
                for candidate in path.rglob("*")
                if candidate.is_file()
                and candidate.suffix.lower() in SUPPORTED_SUFFIXES
                and not any(part.startswith(".") for part in candidate.parts)
            )
    return sorted(set(files))


def load_document(path: Path) -> List[DocumentPart]:
    suffix = path.suffix.lower()
    if suffix == ".md":
        return _load_markdown(path)
    if suffix == ".txt":
        return _load_text(path)
    if suffix == ".pdf":
        return _load_pdf(path)
    if suffix == ".docx":
        return _load_docx(path)
    if suffix == ".xlsx":
        return _load_xlsx(path)
    raise ValueError(f"暂不支持文件格式：{path.suffix or '无扩展名'}")


def _read_text(path: Path) -> str:
    for encoding in ("utf-8", "utf-8-sig", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"无法识别文本编码：{path}")


def _load_text(path: Path) -> List[DocumentPart]:
    text = _read_text(path).strip()
    return [DocumentPart(display_path(path), "正文", text)] if text else []


def _load_markdown(path: Path) -> List[DocumentPart]:
    source = display_path(path)
    sections: List[DocumentPart] = []
    locator = "前言"
    lines: List[str] = []

    def flush() -> None:
        text = "\n".join(lines).strip()
        if text:
            sections.append(DocumentPart(source, locator, text))

    for line in _read_text(path).splitlines():
        heading = line.lstrip().startswith("#") and line.lstrip("#").startswith(" ")
        if heading:
            flush()
            locator = line.lstrip("# ").strip() or "未命名章节"
            lines = [line]
        else:
            lines.append(line)
    flush()
    return sections


def _load_pdf(path: Path) -> List[DocumentPart]:
    from pypdf import PdfReader

    parts: List[DocumentPart] = []
    reader = PdfReader(str(path))
    for page_number, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if text:
            parts.append(DocumentPart(display_path(path), f"第 {page_number} 页", text))
    return parts


def _load_docx(path: Path) -> List[DocumentPart]:
    from docx import Document

    document = Document(str(path))
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_number, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            blocks.append(f"[表格 {table_number}]\n" + "\n".join(rows))
    text = "\n\n".join(blocks)
    return [DocumentPart(display_path(path), "正文", text)] if text else []


def _load_xlsx(path: Path) -> List[DocumentPart]:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
    parts: List[DocumentPart] = []
    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() if value is not None else "" for value in row]
                if any(values):
                    rows.append(" | ".join(values))
            text = "\n".join(rows).strip()
            if text:
                parts.append(DocumentPart(display_path(path), f"工作表：{sheet.title}", text))
    finally:
        workbook.close()
    return parts
